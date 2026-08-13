"""
Output writer — CSV (always) + Excel (if openpyxl is installed).
"""

from __future__ import annotations

import os
from datetime import date
from typing import Optional

import pandas as pd

from _active_config import (
    OUTPUT_COLUMNS,
    OUTPUT_DIR,
)


def write_output(df: pd.DataFrame, run_date: Optional[date] = None,
                 output_dir: str = None, output_columns: list = None) -> None:
    """Write the daily snapshot to CSV and (if openpyxl is available) Excel."""
    if run_date is None:
        run_date = date.today()

    cols = output_columns if output_columns is not None else OUTPUT_COLUMNS
    dest = output_dir    if output_dir    is not None else OUTPUT_DIR

    out = pd.DataFrame()
    for col in cols:
        out[col] = df[col] if col in df.columns else ''

    os.makedirs(dest, exist_ok=True)

    csv_path = os.path.join(dest, f'signals_{run_date.isoformat()}.csv')
    out.to_csv(csv_path, index=False)
    print(f'  [CSV]    Saved → {csv_path}')

    xlsx_path = os.path.join(dest, f'signals_{run_date.isoformat()}.xlsx')
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = str(run_date)
        header_fill = PatternFill(start_color='1a2332', end_color='1a2332', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=10)
        buy_fill  = PatternFill(start_color='d4edda', end_color='d4edda', fill_type='solid')
        sell_fill = PatternFill(start_color='f8d7da', end_color='f8d7da', fill_type='solid')
        thin_border = Border(bottom=Side(style='thin', color='e2e8f0'))
        headers = out.columns.tolist()
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        for row_idx, (_, row) in enumerate(out.iterrows(), 2):
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row[header])
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
            status = str(row.get('confirmation_status', '')).lower()
            fill = buy_fill if 'buy' in status else sell_fill if 'sell' in status else None
            if fill:
                for col_idx in range(1, len(headers) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = fill
        for col_idx, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(str(header)) + 2, 10), 20)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        wb.save(xlsx_path)
        print(f'  [Excel]  Saved → {xlsx_path}')
    except ImportError:
        pass
